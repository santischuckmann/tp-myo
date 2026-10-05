import pytest
from openpyxl import load_workbook

from tp_modelado import resolver_modelo
from tp_modelado.cli import main
from tp_modelado.reporte import a_datos, crear_documento, exportar_excel, guardar_json


@pytest.mark.solver
def test_json_y_excel_tienen_formato_requerido(tmp_path, path_manual):
    resultado = resolver_modelo(str(path_manual), 1, 2)
    ejecucion = {
        "id": "manual-e1-r1", "etapa": 1,
        "instancia": {"path": str(path_manual), "sha256": "x", "seed": None, "restrictiva": False},
        "parametros": {"segundos": 1, "tau": 2}, "repeticion": 1,
        "resultado": a_datos(resultado),
    }
    documento = crear_documento([ejecucion], entorno={})
    json_path = tmp_path / "reporte.json"
    xlsx_path = tmp_path / "reporte.xlsx"
    guardar_json(documento, json_path)
    exportar_excel(documento, xlsx_path, json_path)
    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["Etapa 1", "Etapa 2"]
    assert wb["Etapa 1"].max_column == 5
    assert wb["Etapa 1"].max_row == 2


@pytest.mark.solver
def test_cli_no_sobrescribe_sin_bandera(tmp_path, path_manual):
    salida = tmp_path / "resultado"
    args = ["etapa1", str(path_manual), "--segundos", "1", "--tau", "2", "--salida", str(salida)]
    assert main(args) == 0
    assert main(args) == 2
    assert main(args + ["--sobrescribir"]) == 0


def test_etapa2_cero_puede_reportar_path_no_leido(tmp_path):
    salida = tmp_path / "cero"
    args = ["etapa2", "no_existe.txt", "--segundos", "0", "--salida", str(salida)]
    assert main(args) == 0
    import json

    documento = json.loads((tmp_path / "cero.json").read_text(encoding="utf-8"))
    assert documento["ejecuciones"][0]["instancia"]["sha256"] is None
