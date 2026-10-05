from __future__ import annotations

import hashlib
import json
import platform
import sys
from dataclasses import asdict, is_dataclass
from datetime import datetime
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any, TextIO


ENCABEZADOS = (
    "Instancia",
    "Mejor solución factible",
    "Tau",
    "Garantía de optimalidad",
    "Tiempos",
)


def a_datos(valor: Any) -> Any:
    if is_dataclass(valor):
        return {clave: a_datos(dato) for clave, dato in asdict(valor).items()}
    if isinstance(valor, dict):
        return {str(clave): a_datos(dato) for clave, dato in valor.items()}
    if isinstance(valor, (tuple, list)):
        return [a_datos(dato) for dato in valor]
    return valor


def metadatos_entorno() -> dict[str, Any]:
    paquetes = {}
    for nombre in ("pyscipopt", "pytest", "openpyxl"):
        try:
            paquetes[nombre] = version(nombre)
        except PackageNotFoundError:
            paquetes[nombre] = None
    scip = None
    parametros_scip = {}
    try:
        from pyscipopt import Model

        modelo = Model()
        scip = f"{modelo.getMajorVersion()}.{modelo.getMinorVersion()}.{modelo.getTechVersion()}"
        parametros_scip = {
            "timing/clocktype": modelo.getParam("timing/clocktype"),
            "misc/resetstat": modelo.getParam("misc/resetstat"),
            "timing/rareclockcheck": modelo.getParam("timing/rareclockcheck"),
        }
        liberar = getattr(modelo, "free", None)
        if callable(liberar):
            liberar()
    except Exception:
        pass
    generador = Path("generate_instance_a.py")
    return {
        "python": platform.python_version(),
        "implementacion_python": platform.python_implementation(),
        "sistema": platform.platform(),
        "arquitectura_bits": 64 if sys.maxsize > 2**32 else 32,
        "paquetes": paquetes,
        "scip": scip,
        "parametros_scip": parametros_scip,
        "generador_sha256": hashlib.sha256(generador.read_bytes()).hexdigest() if generador.is_file() else None,
    }


def crear_documento(ejecuciones: list[dict[str, Any]], entorno: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "creado_en": datetime.now().astimezone().isoformat(),
        "entorno": entorno or metadatos_entorno(),
        "ejecuciones": a_datos(ejecuciones),
    }


def guardar_json(documento: dict[str, Any], path: str | Path) -> None:
    destino = Path(path)
    destino.parent.mkdir(parents=True, exist_ok=True)
    texto = json.dumps(documento, ensure_ascii=False, allow_nan=False, indent=2)
    destino.write_text(texto + "\n", encoding="utf-8")


def cargar_json(path: str | Path) -> dict[str, Any]:
    with Path(path).open(encoding="utf-8") as archivo:
        documento = json.load(archivo)
    if documento.get("schema_version") != 1 or not isinstance(documento.get("ejecuciones"), list):
        raise ValueError("JSON de resultados inválido o con versión no soportada")
    return documento


def _texto_seguro_excel(valor: Any) -> Any:
    if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def _fila_excel(ejecucion: dict[str, Any], json_path: Path) -> tuple[Any, ...]:
    resultado = ejecucion["resultado"]
    etapa = int(ejecucion["etapa"])
    solucion = resultado.get("solucion") if etapa == 1 else resultado.get("mejor")
    identificador = f"{ejecucion['id']} | {ejecucion['instancia']['path']}"
    if solucion is None:
        motivo = resultado.get("motivo_parada") or resultado.get("estado")
        texto_solucion = f"sin solución encontrada (estado/motivo: {motivo})"
        tau = None
        garantia = "No, sin solución encontrada"
    else:
        vector = json.dumps(solucion["x"], ensure_ascii=False, separators=(",", ":"))
        sufijo = f"; ganancia={solucion['ganancia_operativa']}; beneficio neto={solucion['beneficio_neto']}"
        texto_solucion = f"x={vector}{sufijo}"
        if len(texto_solucion) > 32767:
            campo = "solucion.x" if etapa == 1 else "mejor.x"
            referencia = f"{json_path.name}#ejecuciones[id={ejecucion['id']}].resultado.{campo}"
            texto_solucion = f"vector completo: {referencia}{sufijo}"
        tau = solucion["tau"]
        garantizada = resultado["garantia_tau"] if etapa == 1 else resultado["garantia_global"]
        garantia = ("Sí" if garantizada else "No") + (", para tau fijo" if etapa == 1 else ", global")
    tiempos = resultado["tiempos"]
    etiqueta = "solicitado al solver" if etapa == 1 else "solicitado global"
    texto_tiempos = (
        f"{etiqueta}={tiempos['solicitado_s']:.6f} s; total={tiempos['total_s']:.6f} s; "
        f"solver={tiempos['solver_s']:.6f} s; exceso={tiempos['exceso_s']:.6f} s"
    )
    return tuple(_texto_seguro_excel(v) for v in (identificador, texto_solucion, tau, garantia, texto_tiempos))


def exportar_excel(documento: dict[str, Any], path: str | Path, json_path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    destino = Path(path)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Etapa 1"
    ws2 = wb.create_sheet("Etapa 2")
    for ws in (ws1, ws2):
        ws.append(ENCABEZADOS)
        for celda in ws[1]:
            celda.font = Font(bold=True)
            celda.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 80
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 28
        ws.column_dimensions["E"].width = 70
    referencia_json = Path(json_path)
    for ejecucion in documento["ejecuciones"]:
        ws = ws1 if int(ejecucion["etapa"]) == 1 else ws2
        ws.append(_fila_excel(ejecucion, referencia_json))
        for celda in ws[ws.max_row]:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(destino)


def imprimir_resultado(resultado, *, stream: TextIO = sys.stdout) -> None:
    datos = a_datos(resultado)
    es_busqueda = "mejor" in datos
    solucion = datos.get("mejor") if es_busqueda else datos.get("solucion")
    print(f"Estado: {datos['estado']}", file=stream)
    if es_busqueda:
        print(f"Motivo: {datos['motivo_parada']}", file=stream)
    if solucion is None:
        print("Mejor solución: sin solución encontrada", file=stream)
    else:
        print(f"x = {solucion['x']}", file=stream)
        print(f"Tau = {solucion['tau']}", file=stream)
        print(f"Ganancia operativa = {solucion['ganancia_operativa']}", file=stream)
        print(f"Beneficio neto = {solucion['beneficio_neto']}", file=stream)
    garantia = datos["garantia_global"] if es_busqueda else datos["garantia_tau"]
    alcance = "global" if es_busqueda else "para tau fijo"
    print(f"Garantía de optimalidad {alcance}: {'sí' if garantia else 'no'}", file=stream)
    if es_busqueda:
        print(f"Taus evaluados/certificados: {datos['taus_evaluados']}/{datos['taus_certificados']}", file=stream)
    tiempos = datos["tiempos"]
    print(
        f"Tiempo solicitado={tiempos['solicitado_s']:.6f} s; "
        f"total={tiempos['total_s']:.6f} s; solver={tiempos['solver_s']:.6f} s; "
        f"exceso={tiempos['exceso_s']:.6f} s",
        file=stream,
    )
